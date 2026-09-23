"""Test smoke de la exportación Excel (`/api/reporte.xlsx`).

Cubre la regresión P0.2: `src/reports/excel_report.py` importa `openpyxl`, que
históricamente no estaba en `requirements.txt` y nunca se importaba en ningún
test. Por eso un deploy con `git reset --hard` + `pip install -r requirements.txt`
devolvía 500 (ImportError) en producción y el CI no lo detectaba.

Este test (a) fuerza el import del módulo (si openpyxl no está instalado, falla
de inmediato) y (b) genera las 6 secciones contra una DB `:memory:` y valida que
cada una produce bytes `.xlsx` válidos y legibles por openpyxl.
"""
import io
import zipfile

import pytest

import openpyxl

from src.reports.excel_report import exportar_excel_bytes


SECCIONES = ["inventario", "leche", "finanzas", "pasturas", "sanidad", "reproduccion"]


def test_openpyxl_importable_y_modulo_presente():
    """Guarda la ruptura de producción: el módulo y su dependencia deben importar."""
    assert openpyxl is not None
    # Re-import defensivo: si excel_report falla al importar, aquí revienta.
    from src.reports import excel_report  # noqa: F401
    assert excel_report is not None


@pytest.mark.parametrize("seccion", SECCIONES)
def test_exportar_excel_bytes_genera_xlsx_valido(db, seccion):
    """Cada sección devuelve bytes que son un .xlsx (zip OOXML) legible."""
    contenido = exportar_excel_bytes(db, seccion=seccion, dias=90)

    assert isinstance(contenido, bytes)
    assert len(contenido) > 0
    # Firma mágica de ZIP (los .xlsx son paquetes OOXML comprimiendos en ZIP).
    assert contenido[:2] == b"PK"
    assert zipfile.is_zipfile(io.BytesIO(contenido))

    # Debe poder abrirse como libro de trabajo y tener al menos una hoja.
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    assert len(wb.sheetnames) >= 1


def test_exportar_excel_seccion_desconocida_caer_en_inventario(db):
    """Una sección no reconocida debe degradar a inventario (no lanzar)."""
    contenido = exportar_excel_bytes(db, seccion="seccion_inexistente")
    assert contenido[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    assert wb.active.title == "Hato Activo"


def test_exportar_excel_incluye_datos_del_hato_activo(db):
    """Con un animal ACTIVO sembrado, el inventario debe reflejarlo y el
    histórico (estado distinto de ACTIVO) debe quedar fuera (regla de oro)."""
    db.conn.execute("INSERT INTO animales (tag, nombre, sexo, raza, estado, fecha_nacimiento) VALUES (?, ?, ?, ?, ?, ?)",
                    ("JJJ001", "Vaca Test", "HEMBRA", "Holstein", "ACTIVO", "2021-01-01"))
    db.conn.execute("INSERT INTO animales (tag, nombre, sexo, raza, estado, fecha_nacimiento) VALUES (?, ?, ?, ?, ?, ?)",
                    ("JJJ999", "Vaca Muerta", "HEMBRA", "Holstein", "MUERTO", "2019-01-01"))
    db.conn.commit()

    contenido = exportar_excel_bytes(db, seccion="inventario")
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    ws = wb.active

    celdas = [str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None]
    assert "JJJ001" in celdas
    assert "JJJ999" not in celdas


@pytest.mark.parametrize("seccion", SECCIONES)
def test_exportar_excel_con_datos_sembrados(db, seccion):
    """Con datos reales en las tablas, cada sección debe renderizar filas sin
    lanzar (ejercita la conversión de tipos y el join de cada consulta)."""
    db.registrar_potrero(nombre="LECHERAS", codigo="01", area_has=3.5,
                         tipo_pasto="Kikuyo")
    db.registrar_animal("JJJ001", nombre="Margarita", sexo="HEMBRA", raza="Holstein",
                        estado="ACTIVO", fecha_nacimiento="2020-01-01", potrero="LECHERAS",
                        hierro="JA", color="Blanco y negro")
    db.registrar_animal("JJJ002", nombre="Toro Padron", sexo="MACHO", raza="Brahman",
                        estado="ACTIVO", fecha_nacimiento="2018-01-01")
    db.registrar_leche("JJJ001", fecha="2026-08-01", litros=12.5, notas="Ordeño AM")
    db.registrar_leche("JJJ001", fecha="2026-08-01", litros=11.0, notas="Ordeño PM")
    db.registrar_finanza(fecha="2026-08-01", tipo="INGRESO", categoria="VENTA_LECHE",
                         concepto="Recibo quincena", monto=400000, litros=23.5,
                         contraparte="Cooperativa", potrero="01")
    db.registrar_tratamiento(animal_tag="JJJ001", fecha="2026-08-05", producto="Oxitetraciclina",
                             dosis="20ml", via="IM", dias_retiro_carne=28, dias_retiro_leche=4,
                             diagnostico="Mastitis")
    db.registrar_servicio(vaca_tag="JJJ001", fecha="2026-08-10", tipo_servicio="IA",
                          toro_pajilla="502")
    db.registrar_diagnostico("JJJ001", fecha="2026-09-01", resultado="PREÑADA")

    contenido = exportar_excel_bytes(db, seccion=seccion, dias=120)
    assert isinstance(contenido, bytes)
    assert contenido[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    assert len(wb.sheetnames) >= 1


def test_exportar_excel_inventario_deriva_categoria_sg(db):
    """La columna ``categoria`` NO existe en ``animales``: la categoría SG se
    deriva en Python por sexo + edad. Un ternero reciente debe clasificarse."""
    # Ternera nacida hoy -> "Hembras <1 año (Ternera)"
    from datetime import date
    db.registrar_animal("T001", nombre="Ternerita", sexo="HEMBRA", raza="Holstein",
                        estado="ACTIVO", fecha_nacimiento=date.today().isoformat())

    contenido = exportar_excel_bytes(db, seccion="inventario")
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    ws = wb.active
    celdas = [str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None]
    assert "T001" in celdas
    assert any("Ternera" in c for c in celdas)

